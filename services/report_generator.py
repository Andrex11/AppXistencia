"""
Generador de reportes en Excel.

Exporta los resultados de la verificación de asistencia a archivos Excel
con formato profesional: colores, encabezados, autoajuste de columnas.
"""

from datetime import date, time
from pathlib import Path
from typing import Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
    numbers,
)
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from services.attendance_engine import (
    PUNTUAL,
    TARDE,
    AUSENTE,
    ANTICIPADO,
    SIN_REGISTRO,
    ERROR,
    NO_APLICA,
)


# ──────────────────────────────────────────────────────────────────────
# Colores por estado
# ──────────────────────────────────────────────────────────────────────
COLOR_PUNTUAL = "FFFFFF"       # Blanco
COLOR_TARDE = "F3F4F6"         # Gris muy suave
COLOR_AUSENTE = "E5E7EB"       # Gris claro neutro
COLOR_ANTICIPADO = "F3F4F6"    # Gris suave
COLOR_SIN_REGISTRO = "F3F4F6"  # Gris suave
COLOR_ERROR = "D1D5DB"         # Gris medio
COLOR_NO_APLICA = "F9FAFB"     # Off-white

COLOR_HEADER_BG = "1F2937"     # Carbón oscuro
COLOR_HEADER_FG = "FFFFFF"     # Blanco

ESTADO_COLORS = {
    PUNTUAL: COLOR_PUNTUAL,
    TARDE: COLOR_TARDE,
    AUSENTE: COLOR_AUSENTE,
    ANTICIPADO: COLOR_ANTICIPADO,
    SIN_REGISTRO: COLOR_SIN_REGISTRO,
    ERROR: COLOR_ERROR,
    NO_APLICA: COLOR_NO_APLICA,
}


class ReportGenerator:
    """Genera reportes de asistencia en formato Excel profesional."""

    def export_to_excel(
        self,
        df: pd.DataFrame,
        filepath: str,
        title: str = "Reporte de Asistencia",
    ) -> str:
        """
        Exporta un DataFrame de resultados a Excel con formato profesional.

        Args:
            df: DataFrame con los resultados de asistencia.
            filepath: Ruta del archivo de salida (.xlsx).
            title: Título del reporte.

        Returns:
            Ruta del archivo generado.
        """
        if df.empty:
            raise ValueError("No hay datos para exportar.")

        # Filtrar columnas internas (que empiezan con _)
        export_cols = [c for c in df.columns if not c.startswith("_")]
        export_df = df[export_cols].copy()

        # Convertir dates y times a strings para Excel
        for col in export_df.columns:
            export_df[col] = export_df[col].apply(self._format_cell_value)

        wb = Workbook()
        ws = wb.active
        ws.title = "Asistencia"

        # Escribir título
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(export_cols))
        title_cell = ws.cell(row=1, column=1, value=title)
        title_cell.font = Font(name="Calibri", size=14, bold=True, color=COLOR_HEADER_FG)
        title_cell.fill = PatternFill(start_color=COLOR_HEADER_BG, end_color=COLOR_HEADER_BG, fill_type="solid")
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 35

        # Fila de fecha de generación
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(export_cols))
        date_cell = ws.cell(row=2, column=1, value=f"Generado: {date.today().strftime('%d/%m/%Y')}")
        date_cell.font = Font(name="Calibri", size=10, italic=True, color="666666")
        date_cell.alignment = Alignment(horizontal="center")

        # Encabezados (fila 4)
        header_row = 4
        header_font = Font(name="Calibri", size=11, bold=True, color=COLOR_HEADER_FG)
        header_fill = PatternFill(start_color=COLOR_HEADER_BG, end_color=COLOR_HEADER_BG, fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        thin_border = Border(
            left=Side(style="thin", color="BDBDBD"),
            right=Side(style="thin", color="BDBDBD"),
            top=Side(style="thin", color="BDBDBD"),
            bottom=Side(style="thin", color="BDBDBD"),
        )

        for col_idx, col_name in enumerate(export_cols, 1):
            cell = ws.cell(row=header_row, column=col_idx, value=col_name)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border

        ws.row_dimensions[header_row].height = 30

        # Datos (desde fila 5)
        data_font = Font(name="Calibri", size=10)
        data_alignment = Alignment(horizontal="center", vertical="center")

        estado_col_idx = None
        if "Estado" in export_cols:
            estado_col_idx = export_cols.index("Estado") + 1

        for row_idx, (_, row_data) in enumerate(export_df.iterrows(), header_row + 1):
            estado_value = None

            for col_idx, col_name in enumerate(export_cols, 1):
                value = row_data[col_name]
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.font = data_font
                cell.alignment = data_alignment
                cell.border = thin_border

                if col_name == "Estado":
                    estado_value = value

            # Aplicar color por estado a toda la fila
            if estado_value:
                fill_color = self._get_color_for_estado(estado_value)
                if fill_color:
                    fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
                    for col_idx in range(1, len(export_cols) + 1):
                        ws.cell(row=row_idx, column=col_idx).fill = fill

        # Autoajuste de columnas
        self._auto_adjust_columns(ws, export_cols, header_row)

        # Agregar filtros automáticos de Excel
        last_row = header_row + len(export_df)
        last_col_letter = get_column_letter(len(export_cols))
        ws.auto_filter.ref = f"A{header_row}:{last_col_letter}{last_row}"

        # Congelar paneles (fijar encabezados)
        ws.freeze_panes = f"A{header_row + 1}"

        # Hoja de resumen
        self._create_summary_sheet(wb, df)

        # Guardar
        output_path = Path(filepath)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(output_path))

        return str(output_path)

    def _format_cell_value(self, value) -> str:
        """Formatea un valor de celda para Excel."""
        if value is None:
            return "—"
        if isinstance(value, date):
            return value.strftime("%d/%m/%Y")
        if isinstance(value, time):
            return value.strftime("%H:%M:%S")
        if isinstance(value, list):
            return ", ".join(str(v) for v in value)
        return str(value)

    def _get_color_for_estado(self, estado_text: str) -> Optional[str]:
        """Obtiene el color de fondo basado en el texto del estado."""
        estado_upper = estado_text.upper()
        for key, color in ESTADO_COLORS.items():
            if key in estado_upper:
                return color
        # Buscar por keywords
        if "PUNTUAL" in estado_upper or "✔" in estado_text:
            return COLOR_PUNTUAL
        if "TARDE" in estado_upper or "⏰" in estado_text:
            return COLOR_TARDE
        if "AUSENTE" in estado_upper or "✖" in estado_text:
            return COLOR_AUSENTE
        if "ANTICIPAD" in estado_upper or "⚡" in estado_text:
            return COLOR_ANTICIPADO
        if "SIN REGISTRO" in estado_upper or "❓" in estado_text:
            return COLOR_SIN_REGISTRO
        if "ERROR" in estado_upper or "⚠" in estado_text:
            return COLOR_ERROR
        return None

    def _auto_adjust_columns(
        self,
        ws: Worksheet,
        columns: list[str],
        header_row: int,
    ) -> None:
        """Autoajusta el ancho de las columnas basándose en el contenido."""
        for col_idx, col_name in enumerate(columns, 1):
            max_length = len(str(col_name))

            # Revisar las primeras 100 filas para determinar el ancho
            for row_idx in range(header_row + 1, min(header_row + 101, ws.max_row + 1)):
                cell = ws.cell(row=row_idx, column=col_idx)
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))

            # Añadir un poco de padding
            adjusted_width = min(max_length + 4, 40)
            ws.column_dimensions[get_column_letter(col_idx)].width = adjusted_width

    def _create_summary_sheet(self, wb: Workbook, df: pd.DataFrame) -> None:
        """Crea una hoja de resumen con estadísticas generales."""
        ws = wb.create_sheet(title="Resumen")

        # Título
        ws.merge_cells("A1:D1")
        title_cell = ws.cell(row=1, column=1, value="Resumen de Asistencia")
        title_cell.font = Font(name="Calibri", size=14, bold=True, color=COLOR_HEADER_FG)
        title_cell.fill = PatternFill(start_color=COLOR_HEADER_BG, end_color=COLOR_HEADER_BG, fill_type="solid")
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 35

        # Estadísticas
        header_font = Font(name="Calibri", size=11, bold=True)
        data_font = Font(name="Calibri", size=11)

        row = 3
        stats = [
            ("Total de registros", len(df)),
            ("Empleados evaluados", df["Empleado"].nunique() if "Empleado" in df.columns else 0),
        ]

        if "Estado" in df.columns:
            estado_col = df["Estado"]
            stats.extend([
                ("", ""),
                ("— Desglose por estado —", ""),
                ("Puntuales", (estado_col == PUNTUAL).sum()),
                ("Llegadas tarde", (estado_col == TARDE).sum()),
                ("Ausentes", (estado_col == AUSENTE).sum()),
                ("Salida anticipada", (estado_col == ANTICIPADO).sum()),
                ("Sin registro", (estado_col == SIN_REGISTRO).sum()),
                ("Errores", (estado_col == ERROR).sum()),
            ])

        for label, value in stats:
            ws.cell(row=row, column=1, value=label).font = header_font
            ws.cell(row=row, column=2, value=value).font = data_font
            row += 1

        # Ajustar anchos
        ws.column_dimensions["A"].width = 30
        ws.column_dimensions["B"].width = 15
